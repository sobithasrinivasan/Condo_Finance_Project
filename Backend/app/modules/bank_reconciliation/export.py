import csv
import io
import json
import logging
from datetime import datetime
from typing import Optional

from .audit import AuditLogger
from .repository import ReconciliationRepository

logger = logging.getLogger(__name__)

SECTION_STATUS_MAP = {
    "matched": ["Matched"],
    "unmatched": ["Suggested", "Unmatched"],
    "manuallyResolved": ["Matched"],  # kept for backward compat with frontend
}

CSV_HEADERS = [
    "ID",
    "Transaction Date",
    "Description",
    "Amount",
    "Type",
    "Reconciliation Type",
    "Matched Record",
    "Matched Description",
    "Match Score",
    "Payment Status",
    "Status",
    "Resolution Notes",
    "Reconciled At",
]

AUDIT_HEADERS = [
    "Reconciliation ID",
    "Timestamp",
    "Action",
    "Performed By",
    "Details",
]


class ExportService:

    def __init__(self, db):
        self.db = db
        self.repo = ReconciliationRepository(db)
        self.audit = AuditLogger(db)

    def _get_records(self, sections: list[str], bank_statement_ids: Optional[list[int]] = None) -> list[dict]:
        statuses = []
        for section in sections:
            statuses.extend(SECTION_STATUS_MAP.get(section, []))

        if not statuses:
            return []

        all_records = []
        for status in statuses:
            rows, _ = self.repo.get_filtered(
                bank_statement_ids=bank_statement_ids,
                status=status,
                page=1,
                page_size=1000,
            )
            all_records.extend(rows)

        return all_records

    def _get_audit_records(self, reconciliation_ids: list[int]) -> list[dict]:
        all_audit = []
        for rec_id in reconciliation_ids:
            trail = self.audit.get_audit_trail("reconciliation", rec_id)
            for entry in trail:
                entry["reconciliation_id"] = rec_id
            all_audit.extend(trail)

        # Also get bank_transaction audit entries
        for rec_id in reconciliation_ids:
            record = self.repo.get_by_id(rec_id, active_only=False)
            if record:
                txn_trail = self.audit.get_audit_trail("bank_transaction", record["bank_transaction_id"])
                for entry in txn_trail:
                    entry["reconciliation_id"] = rec_id
                all_audit.extend(txn_trail)

        all_audit.sort(key=lambda x: x.get("performed_at", ""))
        return all_audit

    def _format_row(self, record: dict) -> list[str]:
        return [
            str(record.get("id", "")),
            str(record.get("transaction_date", "")),
            record.get("transaction_description", ""),
            f"${abs(float(record.get('transaction_amount', 0))):.2f}",
            record.get("transaction_type", ""),
            record.get("reconciliation_type", ""),
            record.get("matched_record_name", ""),
            record.get("matched_record_description", ""),
            str(record.get("match_score", "")),
            record.get("payment_status", ""),
            record.get("status", ""),
            (record.get("resolution_notes", "") or "")[:200],
            str(record.get("matched_date", "")),
        ]

    def _format_audit_row(self, entry: dict) -> list[str]:
        return [
            str(entry.get("reconciliation_id", "")),
            str(entry.get("performed_at", "")),
            entry.get("action", ""),
            entry.get("performed_by_name", "System"),
            entry.get("notes", ""),
        ]

    def export_csv(
        self,
        sections: list[str],
        bank_statement_ids: Optional[list[int]] = None,
        include_audit: bool = False,
    ) -> str:
        output = io.StringIO()
        writer = csv.writer(output)

        # Header
        writer.writerow(["RECONCILIATION REPORT"])
        writer.writerow([f"Generated: {datetime.utcnow().strftime('%Y-%m-%d %H:%M:%S UTC')}"])
        writer.writerow([f"Sections: {', '.join(sections)}"])
        if bank_statement_ids:
            writer.writerow([f"Bank Statement IDs: {', '.join(str(i) for i in bank_statement_ids)}"])
        writer.writerow([])

        records = self._get_records(sections, bank_statement_ids)

        # Reconciliation Records
        writer.writerow(["--- RECONCILIATION RECORDS ---"])
        writer.writerow(CSV_HEADERS)

        for record in records:
            writer.writerow(self._format_row(record))

        writer.writerow([])
        writer.writerow([f"Total Records: {len(records)}"])

        # Summary by status
        writer.writerow([])
        writer.writerow(["--- SUMMARY ---"])
        status_counts = {}
        for r in records:
            s = r.get("status", "Unknown")
            status_counts[s] = status_counts.get(s, 0) + 1
        for s, count in status_counts.items():
            writer.writerow([s, count])

        # Audit History
        if include_audit and records:
            rec_ids = [r["id"] for r in records]
            audit_records = self._get_audit_records(rec_ids)

            writer.writerow([])
            writer.writerow(["--- AUDIT HISTORY ---"])
            writer.writerow(AUDIT_HEADERS)

            for entry in audit_records:
                writer.writerow(self._format_audit_row(entry))

            writer.writerow([])
            writer.writerow([f"Total Audit Entries: {len(audit_records)}"])

        return output.getvalue()

    def export_excel(
        self,
        sections: list[str],
        bank_statement_ids: Optional[list[int]] = None,
        include_audit: bool = False,
    ) -> bytes:
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment
        except ImportError:
            raise ImportError("openpyxl is required for Excel export. Install it with: pip install openpyxl")

        wb = openpyxl.Workbook()

        # Reconciliation Sheet
        ws = wb.active
        ws.title = "Reconciliation"

        header_font = Font(bold=True, color="FFFFFF", size=11)
        header_fill = PatternFill(start_color="0B46AD", end_color="0B46AD", fill_type="solid")

        for col_idx, header in enumerate(CSV_HEADERS, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")

        records = self._get_records(sections, bank_statement_ids)

        for row_idx, record in enumerate(records, 2):
            row_data = self._format_row(record)
            for col_idx, value in enumerate(row_data, 1):
                ws.cell(row=row_idx, column=col_idx, value=value)

        # Auto-width
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 50)

        # Audit Sheet
        if include_audit and records:
            ws_audit = wb.create_sheet("Audit History")

            for col_idx, header in enumerate(AUDIT_HEADERS, 1):
                cell = ws_audit.cell(row=1, column=col_idx, value=header)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center")

            rec_ids = [r["id"] for r in records]
            audit_records = self._get_audit_records(rec_ids)

            for row_idx, entry in enumerate(audit_records, 2):
                row_data = self._format_audit_row(entry)
                for col_idx, value in enumerate(row_data, 1):
                    ws_audit.cell(row=row_idx, column=col_idx, value=value)

            for col in ws_audit.columns:
                max_len = max(len(str(cell.value or "")) for cell in col)
                ws_audit.column_dimensions[col[0].column_letter].width = min(max_len + 2, 50)

        # Summary Sheet
        ws_summary = wb.create_sheet("Summary")
        ws_summary.cell(row=1, column=1, value="Reconciliation Report Summary").font = Font(bold=True, size=14)
        ws_summary.cell(row=2, column=1, value=f"Generated: {datetime.utcnow().strftime('%Y-%m-%d %H:%M:%S UTC')}")
        ws_summary.cell(row=3, column=1, value=f"Sections: {', '.join(sections)}")
        if bank_statement_ids:
            ws_summary.cell(row=4, column=1, value=f"Bank Statement IDs: {', '.join(str(i) for i in bank_statement_ids)}")
        ws_summary.cell(row=5, column=1, value=f"Total Records: {len(records)}")

        status_counts = {}
        for r in records:
            s = r.get("status", "Unknown")
            status_counts[s] = status_counts.get(s, 0) + 1

        row = 6
        ws_summary.cell(row=row, column=1, value="Status").font = Font(bold=True)
        ws_summary.cell(row=row, column=2, value="Count").font = Font(bold=True)
        for s, count in status_counts.items():
            row += 1
            ws_summary.cell(row=row, column=1, value=s)
            ws_summary.cell(row=row, column=2, value=count)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue()

    def export_pdf(
        self,
        sections: list[str],
        bank_statement_ids: Optional[list[int]] = None,
        include_audit: bool = False,
    ) -> bytes:
        try:
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import A4, landscape
            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
        except ImportError:
            raise ImportError("reportlab is required for PDF export. Install it with: pip install reportlab")

        output = io.BytesIO()
        doc = SimpleDocTemplate(output, pagesize=landscape(A4), topMargin=30, bottomMargin=30)
        styles = getSampleStyleSheet()
        elements = []

        # Title
        elements.append(Paragraph("Reconciliation Report", styles["Title"]))
        elements.append(Paragraph(
            f"Generated: {datetime.utcnow().strftime('%Y-%m-%d %H:%M:%S UTC')} | Sections: {', '.join(sections)}",
            styles["Normal"],
        ))
        elements.append(Spacer(1, 20))

        records = self._get_records(sections, bank_statement_ids)

        # Simplified headers for PDF (fewer columns to fit)
        pdf_headers = ["ID", "Date", "Description", "Amount", "Type", "Matched Record", "Score", "Status"]
        table_data = [pdf_headers]

        for record in records:
            table_data.append([
                str(record.get("id", "")),
                str(record.get("transaction_date", "")),
                (record.get("transaction_description", "") or "")[:30],
                f"${abs(float(record.get('transaction_amount', 0))):.2f}",
                record.get("reconciliation_type", ""),
                (record.get("matched_record_name", "") or "")[:25],
                str(record.get("match_score", "")),
                record.get("status", ""),
            ])

        if table_data:
            table = Table(table_data, repeatRows=1)
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0B46AD")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, 0), 8),
                ("FONTSIZE", (0, 1), (-1, -1), 7),
                ("ALIGN", (0, 0), (-1, -1), "LEFT"),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]))
            elements.append(table)

        # Audit section
        if include_audit and records:
            elements.append(Spacer(1, 30))
            elements.append(Paragraph("Audit History", styles["Heading2"]))
            elements.append(Spacer(1, 10))

            rec_ids = [r["id"] for r in records]
            audit_records = self._get_audit_records(rec_ids)

            audit_data = [["Recon ID", "Timestamp", "Action", "By", "Details"]]
            for entry in audit_records:
                audit_data.append([
                    str(entry.get("reconciliation_id", "")),
                    str(entry.get("performed_at", "")),
                    entry.get("action", ""),
                    entry.get("performed_by_name", "System"),
                    (entry.get("notes", "") or "")[:40],
                ])

            if audit_data:
                audit_table = Table(audit_data, repeatRows=1)
                audit_table.setStyle(TableStyle([
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E3A5F")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, -1), 7),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
                    ("TOPPADDING", (0, 0), (-1, -1), 3),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                ]))
                elements.append(audit_table)

        doc.build(elements)
        output.seek(0)
        return output.getvalue()