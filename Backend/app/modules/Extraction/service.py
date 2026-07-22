import os
import uuid
import logging

from fastapi import UploadFile

from app.core.config import settings
from app.modules.extraction.engine import ExtractionEngine
from app.modules.extraction.registry import ExtractorRegistry
from app.modules.extraction.repository import ExtractionRepository
from app.modules.ocr.service import OCRService
from app.prompt.manager import PromptManager

logger = logging.getLogger(__name__)


class ExtractionService:

    def __init__(self, db):

        self.db = db
        self.repo = ExtractionRepository(db)
        try:
            self.ocr = OCRService()
        except Exception as e:
            logger.warning(f"Failed to initialize OCRService: {e}. Multimodal fallback will be used.")
            self.ocr = None
        self.engine = ExtractionEngine()
        self.prompt_manager = PromptManager()

    async def upload_document(
        self,
        file: UploadFile,
        document_type: str,
        source: str = "UPLOAD",
        vendor_id: int | None = None,
        vendor_name: str | None = None,
        uploaded_by: int | None = None,
        document_id: str | None = None
    ):

        file_id = document_id or str(uuid.uuid4())

        os.makedirs(settings.UPLOAD_FOLDER, exist_ok=True)

        filename = f"{file_id}_{file.filename}"
        file_path = os.path.join(
            settings.UPLOAD_FOLDER,
            filename
        )

        with open(file_path, "wb") as buffer:
            buffer.write(await file.read())

        db_id = self.repo.create_document(
            file_id=file_id,
            file_name=file.filename,
            file_path=file_path,
            document_type=document_type,
            source=source,
            status="PROCESSING",
            vendor_id=vendor_id,
            vendor_name=vendor_name,
            uploaded_by=uploaded_by
        )

        return {
            "document_id": file_id,
            "db_id": db_id,
            "status": "PROCESSING"
        }

    def upload_link_document(
        self,
        url: str,
        document_type: str,
        source: str = "EMAIL",
        vendor_id: int | None = None,
        vendor_name: str | None = None,
        uploaded_by: int | None = None,
        document_id: str | None = None
    ):
        import urllib.parse

        file_id = document_id or str(uuid.uuid4())

        # Resolve filename from URL
        parsed_url = urllib.parse.urlparse(url)
        original_filename = os.path.basename(parsed_url.path)
        if not original_filename or "." not in original_filename:
            original_filename = "document.pdf"

        db_id = self.repo.create_document(
            file_id=file_id,
            file_name=original_filename,
            file_path=url,
            document_type=document_type,
            source=source,
            status="PROCESSING",
            vendor_id=vendor_id,
            vendor_name=vendor_name,
            uploaded_by=uploaded_by
        )

        return {
            "document_id": file_id,
            "db_id": db_id,
            "status": "PROCESSING"
        }

    def process_document(
        self,
        document_id: int
    ):

        document = self.repo.get_document(document_id)
        if not document:
            logger.error(f"Document with ID {document_id} not found in database.")
            return None

        file_path = document["file_path"]

        # If it is a URL, download it first
        if file_path.startswith("http://") or file_path.startswith("https://"):
            import requests
            try:
                logger.info(f"Downloading file from URL: {file_path}")
                response = requests.get(file_path, timeout=60)
                response.raise_for_status()

                os.makedirs(settings.UPLOAD_FOLDER, exist_ok=True)
                local_filename = f"{document['document_id']}_{document['document_name']}"
                local_path = os.path.join(settings.UPLOAD_FOLDER, local_filename)

                with open(local_path, "wb") as buffer:
                    buffer.write(response.content)

                # Update path in DB
                cursor = self.db.cursor()
                cursor.execute(
                    "UPDATE document_extraction SET file_path = %s WHERE id = %s",
                    (local_path, document_id)
                )
                self.db.commit()

                document["file_path"] = local_path
                file_path = local_path
                logger.info(f"Downloaded URL content successfully to {local_path}")
            except Exception as dl_err:
                logger.exception(f"Failed to download file from URL {file_path}: {dl_err}")
                self.repo.update_status(document_id, "FAILED", f"File download failed: {str(dl_err)}")
                raise

        extractor = ExtractorRegistry.get_extractor(
            document["document_type"]
        )

        ocr_text = ""
        result = {}

        # Check if Word document (.docx)
        ext = os.path.splitext(file_path)[1].lower()
        is_docx = (ext == ".docx")

        if is_docx:
            logger.info("Extracting text from DOCX Word document...")
            try:
                ocr_text = self._extract_text_from_docx(file_path)
                ocr_text = extractor.pre_process(ocr_text)

                logger.info("Loading prompt template...")
                prompt_template = self.prompt_manager.get_prompt(
                    document_type=document["document_type"]
                )

                logger.info("Running Gemini Extraction...")
                result = self.engine.extract(
                    prompt=prompt_template,
                    ocr_text=ocr_text
                )
            except Exception as docx_err:
                logger.exception(f"DOCX extraction failed: {docx_err}")
                self.repo.update_status(document_id, "FAILED", str(docx_err))
                raise
        else:
            gcp_configured = (
                self.ocr is not None 
                and settings.GCP_PROJECT_ID 
                and (settings.GCP_PROCESSOR_ID or settings.GCP_FORM_PROCESSOR_ID or settings.GCP_LAYOUT_PROCESSOR_ID)
            )

            if gcp_configured:
                try:
                    logger.info("Running OCR...")
                    ocr_text = self.ocr.extract_text(
                        document["file_path"],
                        document_type=document["document_type"]
                    )
                    ocr_text = extractor.pre_process(ocr_text)

                    logger.info("Loading prompt template...")
                    prompt_template = self.prompt_manager.get_prompt(
                        document_type=document["document_type"]
                    )

                    logger.info("Running Gemini Extraction...")
                    result = self.engine.extract(
                        prompt=prompt_template,
                        ocr_text=ocr_text
                    )
                except Exception as ocr_err:
                    logger.warning(f"GCP Document AI or OCR extraction failed, falling back to direct Gemini multimodal extraction: {ocr_err}")
                    gcp_configured = False

            if not gcp_configured:
                logger.info("Running direct Gemini multimodal extraction...")
                try:
                    prompt_template = self.prompt_manager.get_prompt(
                        document_type=document["document_type"]
                    )
                    clean_prompt = prompt_template.replace("{{ocr_text}}", "")

                    mime_type = "application/pdf"
                    if ext == ".png":
                        mime_type = "image/png"
                    elif ext in (".jpg", ".jpeg"):
                        mime_type = "image/jpeg"

                    result = self.engine.extract_from_file(
                        prompt=clean_prompt,
                        file_path=document["file_path"],
                        mime_type=mime_type
                    )
                except Exception as gemini_err:
                    logger.exception(f"Direct Gemini multimodal extraction also failed: {gemini_err}")
                    self.repo.update_status(document_id, "FAILED", str(gemini_err))
                    raise

        result = extractor.post_process(result)

        self.repo.save_result(
            document_id=document_id,
            extracted_json=result,
            ocr_text=ocr_text
        )

        try:
            self._populate_business_tables(document, result)
        except Exception as pop_err:
            logger.exception(f"Error populating business tables: {pop_err}")
            self.repo.update_status(
                document_id,
                "COMPLETED",
                error_message=f"Extraction succeeded but table population failed: {pop_err}"
            )
            return result

        self.repo.update_status(
            document_id,
            "COMPLETED"
        )

        return result

    def _populate_business_tables(self, document: dict, result: dict):
        db_uuid = document["document_id"]
        doc_type = document["document_type"]
        file_path = document["file_path"]
        
        if doc_type.upper() == "BANK_STATEMENT":
            data = result.get("BankStatement", {})
            if not data and "Transactions" in result:
                data = result
                
            info = data.get("Bank_Information", {})
            acc = data.get("Account_Information", {})
            summary = data.get("Account_Summary", {})
            txs = data.get("Transactions", [])
            
            period_str = acc.get("Statement_Period", "")
            month, year = self._parse_period(period_str, txs)
            tx_count = len(txs)
            
            stmt_id = self.repo.create_bank_statement_record(
                document_id=db_uuid,
                file_name=document["document_name"],
                period_month=month,
                period_year=year,
                transaction_count=tx_count,
                file_path=file_path
            )
            
            mapped_txs = []
            for tx in txs:
                tx_date = self._parse_date(tx.get("Date"))
                withdrawal = self._to_float(tx.get("Withdrawal"))
                deposit = self._to_float(tx.get("Deposit"))
                
                amount = 0.0
                tx_type = "Debit"
                if deposit > 0:
                    amount = deposit
                    tx_type = "Credit"
                elif withdrawal > 0:
                    amount = withdrawal
                    tx_type = "Debit"
                else:
                    val = tx.get("Amount")
                    if val is not None:
                        amount = abs(self._to_float(val))
                        tx_type = "Credit" if self._to_float(val) >= 0 else "Debit"
                        
                mapped_txs.append({
                    "transaction_date": tx_date,
                    "description": tx.get("Description", "No Description"),
                    "amount": amount,
                    "type": tx_type
                })
                
            self.repo.create_bank_transaction_records(stmt_id, mapped_txs)
            logger.info(f"Successfully populated bank_statements and bank_transactions for {db_uuid}")
            
        else:
            invoice_data = result.get("Invoice", {})
            if not invoice_data and "Invoice_Information" in result:
                invoice_data = result
                
            v_info = invoice_data.get("Vendor_Information", {})
            inv_info = invoice_data.get("Invoice_Information", {})
            summary = invoice_data.get("Invoice_Summary", {})
            
            vendor_name = document.get("vendor_name")
            if not vendor_name:
                vendor_name = v_info.get("Vendor_Name")
            if not vendor_name:
                vendor_name = "Unknown Vendor"
                
            vendor_id = document.get("vendor_id")
            if not vendor_id:
                vendor_id = self.repo.get_or_create_vendor(vendor_name)
                
            inv_number = inv_info.get("Invoice_Number")
            if not inv_number:
                inv_number = f"INV-{document['document_id'][:8]}"
                
            inv_date = self._parse_date(inv_info.get("Invoice_Date"))
            due_date = self._parse_date(inv_info.get("Due_Date"))
            
            amount = self._to_float(summary.get("Total_Due"))
            if amount == 0.0:
                amount = self._to_float(summary.get("Subtotal"))
                
            notes = invoice_data.get("Additional_Information", {}).get("Notes", "")
            
            self.repo.create_invoice_record(
                document_id=db_uuid,
                vendor_id=vendor_id,
                invoice_number=inv_number,
                amount=amount,
                invoice_date=inv_date,
                due_date=due_date,
                notes=notes,
                file_path=file_path
            )
            logger.info(f"Successfully populated invoice for {db_uuid}")

    def _to_float(self, value) -> float:
        if value is None:
            return 0.0
        if isinstance(value, (int, float)):
            return float(value)
        try:
            clean = "".join(c for c in str(value) if c.isdigit() or c in (".", "-"))
            return float(clean)
        except ValueError:
            return 0.0

    def _parse_date(self, value) -> str:
        import datetime
        if not value:
            return datetime.date.today().strftime("%Y-%m-%d")
        
        val_str = str(value).strip()
        for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%d/%m/%Y", "%Y/%m/%d", "%B %d, %Y", "%b %d, %Y", "%d-%b-%Y", "%d-%B-%Y"):
            try:
                dt = datetime.datetime.strptime(val_str, fmt)
                return dt.strftime("%Y-%m-%d")
            except ValueError:
                continue
                
        try:
            if val_str.isdigit() and len(val_str) > 8:
                dt = datetime.datetime.fromtimestamp(int(val_str[:10]))
                return dt.strftime("%Y-%m-%d")
        except Exception:
            pass
            
        return datetime.date.today().strftime("%Y-%m-%d")

    def _parse_period(self, period_str: str, txs: list) -> tuple[int, int]:
        import datetime
        import re
        
        current_date = datetime.date.today()
        month = current_date.month
        year = current_date.year
        
        if period_str:
            months_map = {
                "january": 1, "february": 2, "march": 3, "april": 4, "may": 5, "june": 6,
                "july": 7, "august": 8, "september": 9, "october": 10, "november": 11, "december": 12,
                "jan": 1, "feb": 2, "mar": 3, "apr": 4, "jun": 6, "jul": 7, "aug": 8, "sep": 9, "oct": 10, "nov": 11, "dec": 12
            }
            for name, m in months_map.items():
                if name in period_str.lower():
                    month = m
                    break
            match = re.search(r"\b(20\d{2})\b", period_str)
            if match:
                year = int(match.group(1))
                return month, year
                
        if txs and txs[0].get("Date"):
            first_tx_date = self._parse_date(txs[0].get("Date"))
            try:
                dt = datetime.datetime.strptime(first_tx_date, "%Y-%m-%d")
                return dt.month, dt.year
            except Exception:
                pass
                
        return month, year

    def _get_db_id(self, document_id) -> int:
        if isinstance(document_id, str) and not str(document_id).isdigit():
            doc = self.repo.get_document_by_uuid(document_id)
            return doc["id"] if doc else 0
        return int(document_id)

    def get_documents(
        self,
        document_id: str | None = None,
        vendor_name: str | None = None,
        doc_type: str | None = None,
        status: str | None = None,
        uploaded_from: str | None = None,
        uploaded_to: str | None = None
    ):
        return self.repo.get_documents_with_filters(
            document_id=document_id,
            vendor_name=vendor_name,
            doc_type=doc_type,
            status=status,
            uploaded_from=uploaded_from,
            uploaded_to=uploaded_to
        )

    def get_document(self, document_id):
        if isinstance(document_id, str) and not str(document_id).isdigit():
            return self.repo.get_document_by_uuid(document_id)
        return self.repo.get_document(int(document_id))

    def get_document_status(self, document_id):
        db_id = self._get_db_id(document_id)
        return self.repo.get_status(db_id)

    def get_extraction_result(self, document_id):
        db_id = self._get_db_id(document_id)
        return self.repo.get_result(db_id)

    def update_document(self, document_id, payload):
        db_id = self._get_db_id(document_id)
        return self.repo.update_document(db_id, payload)

    def delete_document(self, document_id):
        if isinstance(document_id, str) and not str(document_id).isdigit():
            doc = self.repo.get_document_by_uuid(document_id)
        else:
            doc = self.repo.get_document(int(document_id))
            
        if doc:
            file_path = doc.get("file_path")
            if file_path and os.path.exists(file_path):
                try:
                    os.remove(file_path)
                except Exception:
                    pass
            return self.repo.delete_document(doc["id"])
        return {"message": "Document not found."}

    def _extract_text_from_docx(self, file_path: str) -> str:
        import docx
        try:
            doc = docx.Document(file_path)
            full_text = []
            for para in doc.paragraphs:
                full_text.append(para.text)
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        full_text.append(cell.text)
            return '\n'.join(full_text)
        except Exception as e:
            logger.exception(f"Failed to extract text from docx: {e}")
            raise


